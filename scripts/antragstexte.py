#!/usr/bin/env python3
"""
Antragstexte mit Streichungen und Einfügungen
=============================================
Die Excel-Dateien der Abstimmungsergebnisse (data/raw/*.xlsx) setzen den Text
eines Antrags mit Markierungen: rot durchgestrichen ist, was der Antrag
streicht, rot ist, was er einfügt. scripts/scraper.py liest nur den reinen Text
und legt beides ununterscheidbar in `details` ab («längstens jedoch für 14 21
Tage»). Dieses Skript liest die Markierungen nach und schreibt sie getrennt ab:

  data/antragstexte.json
      {"sitzungen": {"<slug>": {"quelle": "<xlsx>", "abstimmungen": {
          "<nr>": {"segmente": [["gleich", "…"], ["alt", "…"], ["neu", "…"]],
                   "hinweise": ["Gilt auch für Art. 71 Abs. 2"]}}}}}

  alt     gestrichen (durchgestrichen)
  neu     eingefügt (rot, nicht durchgestrichen)
  gleich  unverändert

Sitzungen, deren Quelle ein PDF ist, fehlen: dort gibt es keine Markierungen.
Die Blockgrenzen folgen scraper.py (Zeile «Abstimmung N» in der Nummernspalte,
Text in der Spalte rechts daneben, ohne «Ja bedeutet»/«Nein bedeutet»).

Ausführen:
    python3 scripts/antragstexte.py            # alle Sitzungen mit Excel-Quelle
    python3 scripts/antragstexte.py --neueste 4
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from prototyp import sess_sort_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
RAW = DATA / "raw"
ZIEL = DATA / "antragstexte.json"

ABST = re.compile(r"^\s*Abstimmung\s+(\d+)\s*$")
OHNE = ("ja bedeutet", "nein bedeutet", "fakultatives")


def rgb(font):
    """ARGB-Wert der Schriftfarbe oder None (Theme- und Indexfarben: None)."""
    farbe = getattr(font, "color", None)
    wert = getattr(farbe, "rgb", None) if farbe is not None else None
    return wert if isinstance(wert, str) and len(wert) in (6, 8) else None


def ist_rot(font):
    w = rgb(font)
    if not w:
        return False
    w = w[-6:]
    try:
        r, g, b = int(w[0:2], 16), int(w[2:4], 16), int(w[4:6], 16)
    except ValueError:
        return False
    return r >= 0xC0 and g <= 0x60 and b <= 0x60


def art(font):
    if font is not None and getattr(font, "strike", None):
        return "alt"
    if font is not None and ist_rot(font):
        return "neu"
    return "gleich"


def segmente_der_zelle(zelle):
    """Liste (art, text) einer Zelle. Formatierte Läufe aus CellRichText,
    sonst die Zellschrift für den ganzen Text."""
    v = zelle.value
    if v is None:
        return []
    if isinstance(v, CellRichText):
        aus = []
        for teil in v:
            if isinstance(teil, TextBlock):
                aus.append((art(teil.font), teil.text))
            else:                               # Lauf ohne eigene Schrift
                aus.append((art(zelle.font), str(teil)))
        return aus
    return [(art(zelle.font), str(v))]


def zusammenfassen(seg):
    """Benachbarte Läufe gleicher Art verbinden, leere entfernen."""
    aus = []
    for a, t in seg:
        if not t:
            continue
        if aus and aus[-1][0] == a:
            aus[-1][1] += t
        else:
            aus.append([a, t])
    return aus


def spalten(ws):
    """Nummernspalte (mit «Abstimmung N») und Textspalte rechts daneben."""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4):
        for c in row:
            if isinstance(c.value, str) and ABST.match(c.value):
                return c.column, c.column + 1
    return None, None


def datei_lesen(pfad):
    wb = openpyxl.load_workbook(pfad, rich_text=True, data_only=True)
    ws = wb.worksheets[0]
    col_nr, col_tx = spalten(ws)
    if not col_nr:
        return {}
    starts = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col_nr).value
        if isinstance(v, str) and ABST.match(v):
            starts.append((r, int(ABST.match(v).group(1))))
    ergebnis = {}
    for k, (r, nr) in enumerate(starts):
        ende = starts[k + 1][0] if k + 1 < len(starts) else min(ws.max_row + 1, r + 16)
        ende = min(ende, r + 16)
        seg, hinweise = [], []
        for rr in range(r + 1, ende):
            z = ws.cell(rr, col_tx)
            t = z.value
            if t is not None:
                s = str(t).strip()
                if s.startswith("Die Abstimmung"):      # Vorspann des nächsten Geschäfts
                    break
                if s and not s.lower().startswith(OHNE):
                    teile = [t for t in segmente_der_zelle(z) if t[1]]
                    # Zeilenumbruch der Excel-Zelle als Leerschlag; zwischen zwei
                    # Läufen derselben Art gehört er zu diesem Lauf
                    if seg and teile:
                        if seg[-1][0] == teile[0][0]:
                            seg[-1] = (seg[-1][0], seg[-1][1] + " ")
                        else:
                            seg.append(("gleich", " "))
                    seg.extend(teile)
            # rote Hinweise rechts neben dem Text («Gilt auch für Art. …»)
            for cc in range(col_tx + 1, col_tx + 4):
                h = ws.cell(rr, cc)
                if isinstance(h.value, str) and h.value.strip():
                    hs = h.value.strip()
                    if hs.lower().startswith(OHNE) or hs.lower().startswith("zustimmung"):
                        continue
                    if ist_rot(h.font) and re.search(r"\b(Gilt|gilt|Art\.)", hs):
                        hinweise.append(re.sub(r"\s+", " ", hs))
        seg = zusammenfassen(seg)
        # Leerraum normalisieren, Text sonst unverändert
        for s in seg:
            s[1] = re.sub(r"\s+", " ", s[1])
        if seg and seg[0][1].strip() == "":
            seg = seg[1:]
        if not seg and not hinweise:
            continue
        ergebnis[str(nr)] = {"segmente": seg, "hinweise": hinweise}
    return ergebnis


def slug(sitzung):
    teile = sitzung.split("·")
    rest = teile[1].strip() if len(teile) > 1 else ""
    m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})\s*(?:\((.+)\))?", rest)
    if not m:
        return None
    s = f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    if m.group(4):
        s += "-" + re.sub(r"[^a-z]", "", m.group(4).lower())
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--neueste", type=int, help="nur die N neuesten Sitzungen")
    a = ap.parse_args()
    d = json.loads((DATA / "all_sessions.json").read_text(encoding="utf-8"))
    sessions = sorted(d["sessions"], key=sess_sort_key, reverse=True)
    if a.neueste:
        sessions = sessions[:a.neueste]
    alt = {}
    if ZIEL.exists():
        alt = json.loads(ZIEL.read_text(encoding="utf-8")).get("sitzungen", {})
    neu = dict(alt)
    n_dat, n_abst, n_mark = 0, 0, 0
    for s in sessions:
        q = s.get("quelle") or ""
        sl = slug(s["sitzung"])
        if not sl or not q.lower().endswith(".xlsx") or not (RAW / q).exists():
            continue
        try:
            abst = datei_lesen(RAW / q)
        except Exception as e:                       # einzelne Altdatei darf nicht alles stoppen
            print(f"  {q}: {e}")
            continue
        # nur Abstimmungen, die es in der Sitzung gibt
        abst = {k: v for k, v in abst.items() if int(k) <= s["n_votes"]}
        neu[sl] = {"quelle": q, "abstimmungen": abst}
        n_dat += 1
        n_abst += len(abst)
        n_mark += sum(1 for v in abst.values() if any(x[0] != "gleich" for x in v["segmente"]))
    ZIEL.write_text(json.dumps({
        "hinweis": "Antragstexte aus den Excel-Dateien der Abstimmungsergebnisse mit "
                   "Markierungen: alt = gestrichen (durchgestrichen), neu = eingefügt (rot), "
                   "gleich = unverändert. Erzeugt von scripts/antragstexte.py.",
        "sitzungen": dict(sorted(neu.items(), reverse=True)),
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{n_dat} Dateien, {n_abst} Abstimmungen, davon {n_mark} mit Streichung oder Einfügung.")
    print(f"Geschrieben: {ZIEL.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
