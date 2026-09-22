#!/usr/bin/env python3
"""Abgleich: gespeicherte Abstimmungsdaten gegen die Dateien des Kantonsrats.

Prüft für jede Sitzung, deren Quelldatei in data/raw/ liegt, vier Dinge:

1. Stimmen. Jede Einzelstimme und die Summen Ja, Nein, Enthaltung und V/A/N
   aus der Datei gegen data/all_sessions.json, dazu die gedruckte Summenzeile.
2. Hinweistexte. Die Zeilen «Ja bedeutet …» und «Nein bedeutet …» im Anhang der
   Datei gegen das Feld inverted_note.
3. Richtung. Nennt die Zeile «Nein bedeutet Zustimmung Antrag X» denselben
   Antragsteller wie der Titel der Abstimmung, ist die Richtung umgekehrt: ein
   Ja im Rat heisst dann Ablehnung dieses Antrags. Das wird gegen das Feld
   richtung_invertiert gehalten. Diese Prüfung braucht kein Wortprotokoll, sie
   liest nur, was der Kantonsrat selbst gedruckt hat.
4. Probe am Antragsteller. Wer einen Antrag stellt, stimmt ihm zu. Stimmt die
   im Titel genannte Person gegen den eigenen Antrag, ist entweder die Richtung
   falsch oder die Person war nicht der Antragsteller. Nur ein Verdacht, darum
   getrennt ausgewiesen.

    python3 scripts/quellen_pruefen.py                 alle Sitzungen
    python3 scripts/quellen_pruefen.py --seit 2025     ab Jahrgang
    python3 scripts/quellen_pruefen.py --sitzung 21.09.2026

Das Skript ändert nichts. Es schreibt den Bericht auf die Konsole und, mit
--bericht, nach data/quellen_pruefung.md.
"""
import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
RAW = DATA / "raw"

STIMMEN = ("Ja", "Nein", "Enth", "V/A/N")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def txt(ws, r, c):
    if r < 1 or c < 1 or r > ws.max_row or c > ws.max_column:
        return None
    v = ws.cell(r, c).value
    return v if v is not None else None


def lies_quelle(pfad):
    """Stimmen, Summenzeile und Anhangstexte aus einer Abstimmungsdatei."""
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # Kopfzeile mit «Abst. 1», «Abst. 2» … finden.
    kopf_r, spalten = None, []
    for r in range(1, min(ws.max_row, 12) + 1):
        treffer = [c for c in range(1, ws.max_column + 1)
                   if isinstance(txt(ws, r, c), str)
                   and re.match(r"abst\.?\s*\d+", norm(txt(ws, r, c)))]
        if treffer:
            kopf_r, spalten = r, treffer
            break
    if not kopf_r:
        return None

    namen_c = None
    for c in range(1, 6):
        if norm(txt(ws, kopf_r, c)).startswith("nachname"):
            namen_c = c
            break
    namen_c = namen_c or 1

    stimmen, r = {}, kopf_r + 1
    while r <= ws.max_row:
        nach = txt(ws, r, namen_c)
        if not isinstance(nach, str) or not nach.strip():
            break
        vor = txt(ws, r, namen_c + 1)
        schluessel = f"{str(nach).strip()} {str(vor or '').strip()}".strip()
        if norm(nach).startswith("nachname"):      # wiederholte Kopfzeile
            r += 1
            continue
        stimmen[schluessel] = [str(txt(ws, r, c) or "").strip() for c in spalten]
        r += 1

    # Summenzeile: «Ja», «Nein», «Enthaltung», «V / A / N», «Total».
    summen = {}
    for rr in range(r, min(r + 12, ws.max_row) + 1):
        for c in range(1, 6):
            k = norm(txt(ws, rr, c))
            if k in ("ja", "nein", "enthaltung", "v / a / n", "v/a/n", "total"):
                werte = [txt(ws, rr, cc) for cc in spalten]
                if any(isinstance(w, (int, float)) for w in werte):
                    summen[k.replace(" ", "")] = [int(w) if isinstance(w, (int, float)) else None
                                                  for w in werte]
                break

    # Anhang: je Abstimmung Titel, Typ und die beiden Bedeutungszeilen.
    anhang = {}
    for rr in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column, 4) + 1):
            z = txt(ws, rr, c)
            if not isinstance(z, str) or not re.match(r"abstimmung\s+\d+$", norm(z)):
                continue
            nr = int(re.search(r"\d+", z).group())
            e = anhang.setdefault(nr, {"titel": "", "ja": None, "nein": None})
            e["titel"] = str(txt(ws, rr, c + 1) or "").strip()
            for r2 in range(rr, rr + 14):
                stop = any(isinstance(txt(ws, r2, cc), str)
                           and re.match(r"abstimmung\s+\d+$", norm(txt(ws, r2, cc)))
                           for cc in range(1, 4)) and r2 > rr
                if stop:
                    break
                for cc in range(1, ws.max_column + 1):
                    w = txt(ws, r2, cc)
                    if not isinstance(w, str):
                        continue
                    ganz = " ".join(str(txt(ws, r2, c3)).strip()
                                    for c3 in range(cc, min(cc + 6, ws.max_column + 1))
                                    if isinstance(txt(ws, r2, c3), str) and str(txt(ws, r2, c3)).strip())
                    if norm(w).startswith("ja bedeutet") and e["ja"] is None:
                        e["ja"] = re.sub(r"\s+", " ", ganz).strip()
                    if norm(w).startswith("nein bedeutet") and e["nein"] is None:
                        e["nein"] = re.sub(r"\s+", " ", ganz).strip()
            break
    return {"stimmen": stimmen, "summen": summen, "anhang": anhang,
            "n": len(spalten)}


STOP = {"der", "des", "die", "das", "dem", "den", "und", "von", "vom", "abs",
        "art", "lit", "neu", "wie", "folgt", "auf", "zur", "zum", "fur",
        "ja", "nein", "bedeutet", "zustimmung", "annahme", "unterstutzung",
        "antrag", "antrage", "antrags", "gemass", "sowie"}


def worte(text):
    """Tragende Wörter eines Titels oder Hinweises, klein und ohne Akzente."""
    if not text:
        return set()
    t = re.sub(r"\s*\(.*?\)\s*", " ", norm(text))
    return {w for w in re.split(r"[^a-z0-9]+", t)
            if len(w) > 2 and not any(z.isdigit() for z in w) and w not in STOP}


def passt(a, b):
    """Treffen sich zwei Wortmengen? Abkürzungen zählen als Treffer:
    «Flubacher Rüedl.» und «Flubacher Rüedlinger» meinen dieselbe Person."""
    for x in a:
        for y in b:
            if x == y or (len(x) >= 4 and (x.startswith(y) or y.startswith(x))):
                return True
    return False


def akteur(text):
    """Nachname aus «Antrag M. Freivogel» oder «… Zustimmung Antrag M. Freivogel»."""
    w = [x for x in re.split(r"[^a-z0-9]+", re.sub(r"\s*\(.*?\)\s*", " ", norm(text or "")))
         if len(x) > 2 and x not in STOP and not any(z.isdigit() for z in x)]
    # Der Nachname steht zuletzt: «Antrag Matthias Freivogel» und
    # «Antrag M. Freivogel» sollen denselben Akteur ergeben.
    return w[-1] if w else None


GREMIEN = {"spk", "gpk", "kommission", "buro", "buros", "ratsburo", "regierungsrat",
           "rr", "justizkommission", "gesundheitskommission", "erziehungsrat",
           "geschaftsprufungskommission", "spezialkommission", "fraktion"}


def pruefe_sitzung(sess, quelle):
    """Liste von Befunden (art, nr, text)."""
    befunde = []
    q = lies_quelle(quelle)
    if not q:
        return [("technik", None, f"Kopfzeile «Abst. N» nicht gefunden in {quelle.name}")]

    n_daten = sess["n_votes"]
    if q["n"] < n_daten:
        befunde.append(("stimmen", None,
                        f"Datei hat {q['n']} Abstimmungsspalten, gespeichert sind {n_daten}"))

    # 1. Einzelstimmen und Summen
    daten_stimmen = {f"{m['nachname']} {m['vorname']}".strip(): m["votes"]
                     for m in sess["members"]}
    fehlend = sorted(set(q["stimmen"]) - set(daten_stimmen))
    if fehlend:
        befunde.append(("stimmen", None, "in der Datei, aber nicht gespeichert: "
                        + ", ".join(fehlend[:5])))
    # Spalten, in denen die Datei überhaupt Einzelstimmen führt. Fiel die
    # Abstimmungsanlage aus, steht im Blatt nur ein Vermerk und die Zahlen
    # stammen aus der Handauszählung: dort gibt es nichts zu vergleichen.
    mit_einzel = set()
    for i in range(q["n"]):
        n = sum(1 for w in q["stimmen"].values()
                if i < len(w) and w[i].strip() in ("Ja", "Nein", "Enth", "V/A/N"))
        if n >= 10:
            mit_einzel.add(i)
    if q["n"] and not mit_einzel:
        befunde.append(("technik", None,
                        "Datei führt keine Einzelstimmen (Ausfall der Anlage), nicht verglichen"))

    leer = lambda x: str(x or "").replace(" ", "") in ("", "None")
    abweichung = 0
    for name, werte in q["stimmen"].items():
        gespeichert = daten_stimmen.get(name)
        if not gespeichert:
            continue
        for i in range(min(len(werte), n_daten, len(gespeichert))):
            if i not in mit_einzel:
                continue
            a, b = werte[i].replace(" ", ""), str(gespeichert[i]).replace(" ", "")
            if leer(a) and leer(b):
                continue
            if a != b:
                abweichung += 1
                if abweichung <= 5:
                    befunde.append(("stimmen", i + 1,
                                    f"{name}: Datei «{werte[i]}», gespeichert «{gespeichert[i]}»"))
    if abweichung > 5:
        befunde.append(("stimmen", None, f"und {abweichung - 5} weitere Einzelabweichungen"))

    for i in range(n_daten):
        if i not in mit_einzel:
            continue
        gez = {"Ja": 0, "Nein": 0, "Enth": 0, "V/A/N": 0}
        for m in sess["members"]:
            w = m["votes"][i] if i < len(m["votes"]) else "V/A/N"
            gez[w if w in gez else "V/A/N"] += 1
        paare = (("ja", "Ja"), ("nein", "Nein"), ("enthaltung", "Enth"), ("v/a/n", "V/A/N"))
        for schl, k in paare:
            reihe = q["summen"].get(schl.replace(" ", ""))
            if not reihe or i >= len(reihe) or reihe[i] is None:
                continue
            if reihe[i] != gez[k]:
                befunde.append(("summen", i + 1,
                                f"{k}: gedruckt {reihe[i]}, gezählt {gez[k]}"))

    # 2. bis 4. Hinweistexte, Richtung, Probe am Antragsteller
    for v in sess["votes"]:
        nr = int(v["nr"])
        e = q["anhang"].get(nr)
        if not e:
            continue
        ja_datei = re.sub(r"\s+", " ", e["ja"] or "").strip() or None
        ja_daten = re.sub(r"\s+", " ", v.get("inverted_note") or "").strip() or None
        if ja_datei and not ja_daten:
            befunde.append(("hinweis", nr, f"Datei hat «{ja_datei}», gespeichert ist nichts"))
        elif ja_daten and not ja_datei:
            befunde.append(("hinweis", nr, f"gespeichert «{ja_daten}», Datei hat keinen Hinweis"))
        elif ja_datei and ja_daten and norm(ja_datei) != norm(ja_daten):
            befunde.append(("hinweis", nr, f"Datei «{ja_datei}» ≠ gespeichert «{ja_daten}»"))

        inv = bool(v.get("richtung_invertiert"))
        # Ausmehrungen stellen zwei Anträge gegeneinander. Dort gibt es keine
        # inhaltliche Richtung, und offen gelassene Fälle stehen auf der
        # Prüfliste: beides ist keine Abweichung.
        ausmehrung = "ausmehrung" in norm(f"{v.get('typ')} {v.get('titel')}")
        if ausmehrung or v.get("richtung_invertiert") is None:
            w_titel = set()
        else:
            w_titel = worte(v.get("titel"))
        w_ja, w_nein = worte(e["ja"]), worte(e["nein"])
        soll, beleg = None, ""
        if w_titel and w_nein and passt(w_titel, w_nein):
            soll, beleg = True, e["nein"]       # Nein meint den Antrag im Titel
        elif w_titel and w_ja and passt(w_titel, w_ja):
            soll, beleg = False, e["ja"]        # Ja meint den Antrag im Titel
        elif w_titel and w_ja and not w_nein:
            soll, beleg = True, e["ja"]         # Ja meint etwas anderes
        if soll is not None and soll != inv:
            befunde.append(("richtung", nr,
                            f"Datei: «{beleg}», Titel «{v.get('titel')}» "
                            f"→ Umkehr {'ja' if soll else 'nein'}, "
                            f"gespeichert {'ja' if inv else 'nein'}"))
        a_titel = akteur(v.get("titel"))

        # Probe am Antragsteller
        if a_titel and a_titel not in GREMIEN:
            i = nr - 1
            treffer = [m for m in sess["members"] if norm(m["nachname"]) == a_titel]
            if len(treffer) == 1 and 0 <= i < len(treffer[0]["votes"]):
                roh = treffer[0]["votes"][i]
                if roh in ("Ja", "Nein"):
                    dafuer = (roh == "Nein") if inv else (roh == "Ja")
                    if not dafuer:
                        befunde.append(("probe", nr,
                                        f"{treffer[0]['nachname']} stimmte «{roh}» "
                                        f"zum eigenen Antrag, Umkehr {'ja' if inv else 'nein'}"
                                        + (f", Hinweis «{e['ja']}»" if e["ja"] else ", ohne Hinweis")))
    return befunde


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seit", type=int, default=0, help="nur Sitzungen ab diesem Jahr")
    ap.add_argument("--sitzung", help="Textstück aus dem Sitzungsnamen")
    ap.add_argument("--bericht", action="store_true", help="zusätzlich nach data/quellen_pruefung.md")
    a = ap.parse_args()

    d = json.loads((DATA / "all_sessions.json").read_text(encoding="utf-8"))
    zeilen, gepruefte, mit_befund = [], 0, 0
    zaehler = {}
    for s in d["sessions"]:
        if a.sitzung and a.sitzung not in s["sitzung"]:
            continue
        jahr = re.search(r"\.(\d{4})", s["sitzung"])
        if a.seit and (not jahr or int(jahr.group(1)) < a.seit):
            continue
        quelle = RAW / (s.get("quelle") or "")
        if not quelle.exists() or quelle.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        gepruefte += 1
        befunde = pruefe_sitzung(s, quelle)
        for art, _, _ in befunde:
            zaehler[art] = zaehler.get(art, 0) + 1
        if befunde:
            mit_befund += 1
            zeilen.append(f"\n## {s['sitzung']}  ({quelle.name})")
            for art, nr, text in befunde:
                zeilen.append(f"- [{art}]{f' Nr. {nr}:' if nr else ''} {text}")

    kopf = (f"Geprüft: {gepruefte} Sitzungen mit Quelldatei, "
            f"{mit_befund} mit Befund. "
            + (", ".join(f"{k}: {v}" for k, v in sorted(zaehler.items())) or "keine Befunde"))
    print(kopf)
    print("\n".join(zeilen))
    if a.bericht:
        ziel = DATA / "quellen_pruefung.md"
        ziel.write_text("# Abgleich mit den Dateien des Kantonsrats\n\n" + kopf + "\n"
                        + "\n".join(zeilen) + "\n", encoding="utf-8")
        print(f"\ngeschrieben: {ziel.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
